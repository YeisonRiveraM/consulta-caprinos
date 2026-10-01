
from __future__ import annotations

import json
import re
import shutil
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
EXCEL_DIR = ROOT / "excel"
OUTPUT_DIR = ROOT / "datos"
ASSETS_DIR = OUTPUT_DIR / "imagenes"

LIFE_FILE = EXCEL_DIR / "Hojas_Vida.xlsx"
MILK_FILE = EXCEL_DIR / "Leche_Cabras.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "subjects.json"


def normalize(value):
    """Normaliza identificaciones sin depender del nombre de la hoja."""
    if value is None:
        return ""

    text = unicodedata.normalize("NFKC", str(value))
    return text.strip().upper()


def json_value(value):
    """Convierte valores de Excel a tipos compatibles con JSON."""
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if isinstance(value, (str, int, float, bool)):
        return value

    return str(value)


def find_field(ws, label):
    """Busca una etiqueta y devuelve el valor de la celda a su derecha."""
    target = normalize(label)

    for row in ws.iter_rows():
        for cell in row:
            if normalize(cell.value) == target:
                return ws.cell(
                    row=cell.row,
                    column=cell.column + 1
                ).value

    return None


def find_subject_name(ws):
    return find_field(ws, "NOMBRE") or ws.title


def find_subject_id(ws):
    return find_field(ws, "N° IDENTIFICACIÓN")


def extract_images(ws, workbook_name, sheet_index):
    """Extrae las imágenes incrustadas de una hoja."""
    images = []

    for image_index, image in enumerate(
        getattr(ws, "_images", []), start=1
    ):
        try:
            image_bytes = image._data()

            extension = image.format or "png"
            extension = re.sub(
                r"[^a-zA-Z0-9]", "", extension
            ).lower() or "png"

            filename = (
                f"{workbook_name}_hoja{sheet_index}_"
                f"{image_index}.{extension}"
            )

            destination = ASSETS_DIR / filename
            destination.write_bytes(image_bytes)

            anchor = getattr(image, "anchor", None)
            cell = getattr(anchor, "_from", None)

            position = ""
            if cell is not None:
                position = (
                    f"{get_column_letter(cell.col + 1)}"
                    f"{cell.row + 1}"
                )

            images.append({
                "src": f"datos/imagenes/{filename}",
                "alt": f"Imagen de {ws.title}",
                "caption": (
                    f"Imagen original de la hoja {ws.title}"
                    + (f", posición {position}" if position else "")
                ),
            })

        except Exception as error:
            print(
                f"ADVERTENCIA: no se pudo extraer una imagen "
                f"de {ws.title}: {error}"
            )

    return images


def extract_charts(ws):
    """
    Conserva información descriptiva de los gráficos.
    Los datos de origen también se conservan en la tabla completa.
    """
    result = []

    for index, chart in enumerate(
        getattr(ws, "_charts", []), start=1
    ):
        chart_info = {
            "label": f"Gráfico {index}",
            "type": type(chart).__name__,
            "title": "",
            "series": [],
        }

        try:
            title = chart.title
            if title is not None:
                chart_info["title"] = str(title)
        except Exception:
            pass

        for series in getattr(chart, "ser", []):
            item = {}

            for key in ("val", "cat", "tx", "xVal", "yVal"):
                reference = getattr(series, key, None)

                if reference is not None:
                    try:
                        item[key] = str(reference)
                    except Exception:
                        item[key] = repr(reference)

            chart_info["series"].append(item)

        result.append(chart_info)

    return result


def extract_full_sheet(ws, workbook_name, sheet_index):
    """
    Representa el rango completo utilizado de la hoja.
    Se incluyen las columnas y filas vacías dentro de ese rango.
    Las fórmulas se conservan como expresiones de Excel.
    """
    columns = [
        get_column_letter(index)
        for index in range(1, ws.max_column + 1)
    ]

    rows = []

    for row_index in range(1, ws.max_row + 1):
        row = []

        for column_index in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_index, column=column_index)

            value = cell.value

            # Conservamos las fórmulas como fórmulas.
            row.append(json_value(value))

        rows.append(row)

    images = extract_images(
        ws, workbook_name, sheet_index
    )

    charts = extract_charts(ws)

    notes = []

    if charts:
        notes.append({
            "label": "Gráficos incrustados",
            "value": json.dumps(
                charts, ensure_ascii=False, indent=2
            ),
        })

    merged_ranges = [
        str(rng) for rng in ws.merged_cells.ranges
    ]

    if merged_ranges:
        notes.append({
            "label": "Celdas combinadas",
            "value": ", ".join(merged_ranges),
        })

    notes.append({
        "label": "Dimensiones de la hoja",
        "value": (
            f"{ws.max_row} filas × {ws.max_column} columnas"
        ),
    })

    notes.append({
        "label": "Libro de origen",
        "value": workbook_name,
    })

    notes.append({
        "label": "Nombre original de la hoja",
        "value": ws.title,
    })

    section = {
        "id": f"{workbook_name}-{sheet_index}",
        "title": f"{ws.title} · {workbook_name}",
        "description": (
            "Contenido de la hoja original. "
            "Se conservan las posiciones y celdas vacías "
            "dentro del rango utilizado."
        ),
        "tables": [{
            "title": "Contenido completo de la hoja",
            "columns": columns,
            "rows": rows,
        }],
        "images": images,
        "notes": notes,
    }

    return section


def extract_workbook(path, workbook_name, include_template=False):
    """
    Lee todas las hojas de un libro.
    La plantilla también se conserva como ficha consultable.
    """
    if not path.exists():
        raise FileNotFoundError(
            f"No se encuentra el archivo: {path}"
        )

    # data_only=False conserva las fórmulas originales.
    workbook = load_workbook(
        path,
        data_only=False,
        read_only=False,
    )

    records = []

    for sheet_index, ws in enumerate(workbook.worksheets, start=1):
        is_template = normalize(ws.title) == "PLANTILLA"

        if is_template and not include_template:
            continue

        raw_id = find_subject_id(ws)
        name = find_subject_name(ws)

        record = {
            "id": normalize(raw_id) if raw_id else "",
            "displayName": str(name).strip(),
            "sourceSheet": ws.title,
            "sourceWorkbook": workbook_name,
            "isTemplate": is_template,
            "sections": [
                extract_full_sheet(
                    ws, workbook_name, sheet_index
                )
            ],
        }

        # Campos para mostrar un resumen sin reemplazar
        # ni eliminar el contenido completo de la hoja.
        record["summaryFields"] = []

        for label in (
            "NOMBRE",
            "N° IDENTIFICACIÓN",
            "ESPECIE",
            "RAZA",
            "PROCEDENCIA",
            "FECHA NACIMIENTO",
            "PADRE",
            "MADRE",
            "FINALIDAD",
            "EDAD",
        ):
            value = find_field(ws, label)

            record["summaryFields"].append({
                "label": label,
                "value": json_value(value),
            })

        photo = record["sections"][0]["images"]
        if photo:
            record["photo"] = photo[0]["src"]

        records.append(record)

    workbook.close()
    return records


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)

    # Limpiar únicamente las imágenes generadas anteriormente.
    for file in ASSETS_DIR.iterdir():
        if file.is_file():
            file.unlink()

    print("Leyendo hojas de vida...")
    life_records = extract_workbook(
        LIFE_FILE,
        "Hojas_Vida",
        include_template=True,
    )

    print("Leyendo registros de leche...")
    milk_records = extract_workbook(
        MILK_FILE,
        "Leche_Cabras",
        include_template=True,
    )

    # Agrupar por identificación para evitar asociaciones
    # incorrectas cuando existen identificaciones duplicadas.
    life_by_id = defaultdict(list)

    for record in life_records:
        if record["id"] and not record["isTemplate"]:
            life_by_id[record["id"]].append(record)

    diagnostics = []

    for identifier, records in life_by_id.items():
        if len(records) > 1:
            names = ", ".join(
                record["displayName"] for record in records
            )

            diagnostics.append({
                "type": "identificacion_duplicada",
                "id": identifier,
                "message": (
                    f"La identificación {identifier} aparece "
                    f"en varias hojas de vida: {names}. "
                    "No se combinarán automáticamente."
                ),
            })

    # Incorporar los registros de leche al sujeto correcto
    # únicamente cuando la identificación permite una
    # correspondencia inequívoca.
    unmatched_milk = []

    for milk_record in milk_records:
        if milk_record["isTemplate"]:
            continue

        identifier = milk_record["id"]
        matches = life_by_id.get(identifier, [])

        if identifier and len(matches) == 1:
            matches[0]["sections"].extend(
                milk_record["sections"]
            )

            # La fotografía de leche se conserva en su sección.
            matches[0].setdefault(
                "sourceSheets", []
            ).append({
                "workbook": "Leche_Cabras",
                "sheet": milk_record["sourceSheet"],
            })

        else:
            if len(matches) > 1:
                diagnostics.append({
                    "type": "leche_sin_asociacion",
                    "id": identifier,
                    "sheet": milk_record["sourceSheet"],
                    "message": (
                        "La hoja de leche no se ha asociado "
                        "porque la identificación coincide con "
                        "varias hojas de vida."
                    ),
                })

            unmatched_milk.append(milk_record)

    # Mantener accesible cualquier hoja de leche que no
    # haya podido vincularse con una única hoja de vida.
    records = life_records + unmatched_milk

    # Agregar diagnósticos para que sean visibles en los datos
    # generados y no se pierdan silenciosamente.
    payload = {
        "generatedAt": datetime.now().astimezone().isoformat(),
        "subjects": records,
        "diagnostics": diagnostics,
        "sourceFiles": [
            LIFE_FILE.name,
            MILK_FILE.name,
        ],
        "notes": [
            "Las fórmulas se conservan como expresiones.",
            "Los resultados calculados pueden requerir Excel "
            "o LibreOffice para recalcularse.",
            "Los gráficos se documentan mediante metadatos; "
            "los datos de sus hojas de origen se conservan "
            "en las tablas completas.",
            "El rango completo de cada hoja se conserva, "
            "incluyendo las celdas vacías dentro del rango usado.",
        ],
    }

    OUTPUT_FILE.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"Fichas generadas: {len(records)}")
    print(f"Advertencias: {len(diagnostics)}")
    print(f"Archivo generado: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
